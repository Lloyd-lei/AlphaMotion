"""Build the IROS deck as a native PPTX (16:9, 1920x1080 px grid), embedding the real mp4 clips."""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
S = '/tmp/claude-1000/-home-arenalabs-Desktop--be-water--robot-/eff459e5-3969-4f6d-9ec8-eae021b0a07a/scratchpad/iros_deck'
M, W = f'{S}/media', f'{S}/site/up'
COVER = '/home/arenalabs/Desktop/"be water, robot"/alpha_motion2/pitch_deck_20260722/cover_robot.png'
OUT = f'{S}/pptx/AlphaMotion_IROS2026.pptx'
HEAD, BODY, MONO = 'Arial', 'Arial', 'Courier New'
BG, BG2, TXT, SOFT, MUTE, DIM, ORANGE = '0B0B0A', '151513', 'F2F0EA', 'D8D5CC', 'B9B6AD', '6E6C66', 'FF7A00'
px = lambda v: Emu(int(round(v * 6350)))
prs = Presentation(); prs.slide_width, prs.slide_height = px(1920), px(1080)
blank = prs.slide_layouts[6]

def slide(bg):
    s = prs.slides.add_slide(blank); f = s.background.fill; f.solid(); f.fore_color.rgb = RGBColor.from_string(bg); return s

def text(s, x, y, w, h, paras, size, color=TXT, font=BODY, bold=False, align=PP_ALIGN.LEFT, spacing=1.15, anchor=MSO_ANCHOR.TOP, caps=False, tracking=None):
    tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = anchor
    if isinstance(paras, str): paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.line_spacing = spacing
        runs = para if isinstance(para, list) else [(para, color)]
        for t, c in runs:
            r = p.add_run(); r.text = t.upper() if caps else t; f = r.font; f.size = Pt(size / 2); f.name = font; f.bold = bold; f.color.rgb = RGBColor.from_string(c)
            if tracking:
                r._r.get_or_add_rPr().set('spc', str(int(tracking * 100)))
    return tb

def rect(s, x, y, w, h, fill=None, line=None, lw=1, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = px(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def fit(path, x, y, w, h):
    iw, ih = Image.open(path).size; sc = min(w / iw, h / ih); fw, fh = iw * sc, ih * sc
    return x + (w - fw) / 2, y + (h - fh) / 2, fw, fh

def image(s, path, x, y, w, h, backdrop=None):
    if backdrop: rect(s, x, y, w, h, fill=backdrop)
    fx, fy, fw, fh = fit(path, x, y, w, h); return s.shapes.add_picture(path, px(fx), px(fy), px(fw), px(fh))

def video(s, stem, x, y, w, h, backdrop='000000'):
    mp4, jpg = f'{stem}.mp4', f'{stem}.jpg'
    rect(s, x, y, w, h, fill=backdrop)
    fx, fy, fw, fh = fit(jpg, x, y, w, h)
    return s.shapes.add_movie(mp4, px(fx), px(fy), px(fw), px(fh), poster_frame_image=jpg, mime_type='video/mp4')

def eyebrow(s, t, color=ORANGE): text(s, 128, 128, 1664, 36, t, 24, color, MONO, caps=True, tracking=4)
def title(s, t, y=176, h=80, size=64, color=TXT, w=1664): text(s, 128, y, w, h, t, size, color, HEAD, bold=True, spacing=1.05)
def footer(s, t, color=DIM): text(s, 128, 992, 1664, 36, t, 24, color)
def notes(s, t): s.notes_slide.notes_text_frame.text = t
def stat(s, x, y, w, big, small, color=TXT, size=72):
    text(s, x, y, w, size + 8, big, size, color, HEAD, bold=True, spacing=1.0); text(s, x, y + size + 14, w, 40, small, 26, MUTE)

import sys; sys.path.insert(0, f'{S}/v5'); from notes import NOTES
AUTH = 'Lloyd Lei · CMU · IROS 2026'
BLK = '14130F'

# cover
s = slide(BG)
im = Image.open(COVER); iw, ih = im.size; box_w, box_h = 820, 1080
pic = s.shapes.add_picture(COVER, px(1100), px(0), px(box_w), px(box_h))
ar, br = iw / ih, box_w / box_h
if ar > br: c = (1 - br / ar) / 2; pic.crop_left = pic.crop_right = c
else: c = (1 - ar / br) / 2; pic.crop_top = pic.crop_bottom = c
text(s, 128, 330, 900, 36, 'IROS 2026', 24, ORANGE, MONO, tracking=4)
text(s, 128, 380, 980, 220, 'A Cross-Embodiment Motion Tokenizer', 96, TXT, HEAD, bold=True, spacing=1.0)
text(s, 128, 612, 900, 60, 'Learning from joint-to-joint moving synergy', 40, 'C9C6BD')
rect(s, 128, 700, 120, 3, fill=ORANGE)
text(s, 128, 730, 900, 40, 'Lloyd Lei · Carnegie Mellon University', 28, '9C9A93')
text(s, 128, 940, 950, 50, '“A brilliant idea is a secret that no one else knows.” — after Peter Thiel', 28, DIM)
notes(s, NOTES['cover'])

# field
s = slide(BG); eyebrow(s, 'Embodied AI'); title(s, 'The peak of engineering in human history', size=72, h=90)
x = 128
for c in ['Mechanical engineering', 'Deep learning', 'Electromagnetism', 'Materials science', 'Computer engineering']:
    w = len(c) * 15.5 + 56
    rect(s, x, 330, w, 64, line='3A3934', lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.5
    text(s, x, 330, w, 64, c, 28, SOFT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE); x += w + 16
text(s, 128, 520, 520, 180, '2022', 160, ORANGE, HEAD, bold=True, spacing=1.0)
text(s, 700, 530, 1000, 70, 'Neural networks will replace MPC.', 48, TXT, HEAD, bold=True)
text(s, 700, 610, 1000, 120, 'Model-based control gives way to control learned end to end. Four years later, we are on the eve of the next explosive cycle.', 30, MUTE, spacing=1.35)
footer(s, AUTH); notes(s, NOTES['field'])

# starbot
s = slide(BG); eyebrow(s, 'StarBot · World Humanoid Robot Games, Beijing 2025'); title(s, 'Fully autonomous, one year ago')
for i, (b, sm, col) in enumerate([('4th', 'hotel reception', ORANGE), ('6th', 'medicine sorting', TXT), ('1 of 2', 'fully autonomous teams on site', TXT)]):
    stat(s, 128 + i * 520, 300, 500, b, sm, col, 120)
text(s, 128, 590, 1600, 40, 'One year, three paradigm shifts', 28, '9C9A93')
cw = (1664 - 72) / 4
for i, (y, t, col) in enumerate([('2025', 'CV detectors + rule-based control', '3A3934'), ('', 'VLA', '3A3934'), ('', 'WAM', '3A3934'), ('2026', 'Robotics Agent Harness', ORANGE)]):
    x = 128 + i * (cw + 24); rect(s, x, 660, cw, 3, fill=col)
    text(s, x, 683, cw, 30, y, 22, DIM, MONO); text(s, x, 723, cw, 100, t, 34, TXT, HEAD, bold=True, spacing=1.15)
footer(s, 'Founding member, StarBot · early autonomous robotics stack'); notes(s, NOTES['starbot'])

# wbc
s = slide(BG); eyebrow(s, 'Whole-body control'); title(s, 'The control layer became a network', size=72, h=90)
cw = (1664 - 64) / 3
for i, (a, b, col) in enumerate([('Atlas · RAI Institute', 'The whole humanoid control layer, handed to neural networks, running autonomously.', ORANGE),
                                 ('SONIC · NVIDIA GEAR', 'A generalist whole-body controller. The current bar for WBC.', '3A3934'),
                                 ('BFM-Zero · LeCAR Lab', 'Behavior foundation models: a new direction for reinforcement learning.', '3A3934')]):
    x = 128 + i * (cw + 32); rect(s, x, 330, cw, 3, fill=col)
    text(s, x, 361, cw, 50, a, 36, TXT, HEAD, bold=True); text(s, x, 425, cw, 180, b, 28, MUTE, spacing=1.35)
text(s, 128, 700, 1500, 110, 'Yet most commercial stacks are still tracking-based RL: one expert policy per motion.', 34, TXT, spacing=1.35)
footer(s, AUTH); notes(s, NOTES['wbc'])

# palo alto
s = slide(BG); eyebrow(s, 'Where humanoids stand')
text(s, 128, 280, 700, 200, '$40B+', 200, ORANGE, HEAD, bold=True, spacing=0.95)
text(s, 128, 490, 640, 80, 'invested in US robotics, as of Q3 2026', 30, MUTE)
text(s, 860, 270, 930, 260, '“…and yet you still don’t see a single working robot on the streets of Palo Alto.”', 52, TXT, HEAD, bold=True, spacing=1.2)
text(s, 860, 545, 930, 40, '— an investor on Sand Hill Road', 28, '9C9A93')
text(s, 128, 700, 1600, 110, 'Without a breakthrough in low-level control, the public still sees humanoids as expensive, backflipping toys on a gala stage.', 32, SOFT, spacing=1.35)
footer(s, AUTH); notes(s, NOTES['palo-alto'])

# problem
s = slide(BG); eyebrow(s, 'The question I kept asking')
for i, q in enumerate(['Why was SONIC built only on G1, and why has almost no team reproduced it on other hardware?',
                       'Manipulation data is abundant. Why no jaw-dropping breakthrough in humanoid manipulation?']):
    y = 200 + i * 150; rect(s, 128, y, 3, 110, fill='3A3934'); text(s, 160, y, 1500, 120, q, 36, MUTE, spacing=1.35)
text(s, 128, 540, 1600, 40, 'Every answer pointed to one core:', 30, '9C9A93')
text(s, 128, 600, 1664, 280, [[('Robots are no longer starved for data. ', TXT), ('Usable locomotion and manipulation data for a non-mainstream body is still out of reach.', ORANGE)]], 60, font=HEAD, bold=True, spacing=1.1)
footer(s, AUTH); notes(s, NOTES['problem'])

# fuel
s = slide(BG); eyebrow(s, 'May 2026 · fuel for the scaling law'); title(s, 'A learning-based retargeting tokenizer')
video(s, f'{W}/8robots-run', 128, 288, 1110, 624)
for i, t in enumerate(['Zero-shot on new embodiments', 'Zero-shot on new motions', 'Locomotion and manipulation']):
    y = 300 + i * 150; rect(s, 1300, y, 492, 3, fill=ORANGE); text(s, 1300, y + 19, 492, 90, t, 36, TXT, HEAD, bold=True, spacing=1.15)
text(s, 1300, 780, 492, 120, 'One motion from video, eight humanoids at once. AgiBot X2 was never in training.', 26, MUTE, spacing=1.35)
footer(s, 'Video → GenMo (SMPL) → tokenizer → Unitree G1, H1-2 · PND Adam · Booster T1 · Fourier N1, GR-3 · Kuavo S45 · AgiBot X2')
notes(s, NOTES['fuel'])

# how
s = slide(BG2); eyebrow(s, 'How it works'); title(s, 'The body is an input, not a hyperparameter')
steps = [('01 · Any source body', 'A named joint tree from its MJCF: axes, limits, rest frames, plus a language embedding of each joint name.', '33312C'),
         ('02 · Graph encoder', 'A transformer over joints with a learned hop-distance bias: joint-to-joint synergy along the kinematic tree.', '33312C'),
         ('03 · Discrete token', '132 slots per frame, finite scalar quantization. One token per frame, shared by every body.', ORANGE),
         ('04 · Any target body', "A decoder conditioned on the target's descriptor: joint rotations, root height, contacts.", '33312C')]
bw = (1664 - 3 * 64) / 4
for i, (h3, p, lc) in enumerate(steps):
    x = 128 + i * (bw + 64); rect(s, x, 320, bw, 400, fill='1E1D1A', line=lc, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.04
    text(s, x + 28, 352, bw - 56, 40, h3, 24, ORANGE, MONO); text(s, x + 28, 410, bw - 56, 290, p, 28, SOFT, spacing=1.35)
    if i < 3: rect(s, x + bw + 14, 505, 36, 30, fill='6E6C66', shape=MSO_SHAPE.RIGHT_ARROW)
text(s, 128, 770, 1600, 50, 'Joint names are language: a new robot is described the way a person would describe it.', 28, '9C9A93')
footer(s, AUTH); notes(s, NOTES['how'])

# sources
s = slide(BG); eyebrow(s, 'What this means · 1'); title(s, 'Every legacy data pool maps to any new robot')
for i, (st, lab) in enumerate([('source-robot', 'Real-robot teleop logs'), ('source-ego', 'Egocentric capture'), ('source-web', 'Internet and game video')]):
    x = 128 + i * (533 + 32); video(s, f'{W}/{st}', x, 300, 533, 300); text(s, x, 620, 533, 50, lab, 32, TXT, HEAD, bold=True)
text(s, 128, 720, 1600, 110, 'Years of Unitree G1 data, egocentric and UMI capture, real-world teleop: all enter the same token, and decode onto any robot.', 32, SOFT, spacing=1.35)
footer(s, AUTH); notes(s, NOTES['sources'])

# manip
s = slide(BG); eyebrow(s, 'What this means · 2'); title(s, 'Locomotion was the objective. Manipulation emerged.', h=150)
video(s, f'{W}/alphamotion-manipulation-eval', 128, 340, 330, 456); video(s, f'{W}/umi', 490, 340, 810, 456)
stat(s, 1332, 350, 460, '0.03 cm', 'hand SE(3) error', ORANGE, 88); stat(s, 1332, 530, 460, '0 cm', 'floor penetration, feet planted', TXT, 88)
text(s, 128, 825, 1600, 50, 'Left: manipulation evaluation. Right: UMI portable capture, hands only, completed into whole-body motion.', 28, '9C9A93')
footer(s, 'Sparse hand targets → tokenizer completes the body → causal constrained projection, 6–9 ms/frame')
notes(s, NOTES['manip'])

# scaling
s = slide('F4F0E7'); eyebrow(s, 'What this means · 3', 'C25E00')
title(s, 'Usable data per robot grows by orders of magnitude, and error keeps falling', h=150, color=BLK)
image(s, f'{W}/scaling_cross.png', 128, 372, 600, 478); image(s, f'{W}/scaling_unseen.png', 752, 372, 600, 478)
text(s, 128, 868, 600, 40, [[('Cross-embodiment: ', '3D3B35'), ('13.5° → 4.8°', BLK)]], 26)
text(s, 752, 868, 600, 40, [[('Unseen embodiment: ', '3D3B35'), ('15.8° → 6.4°', BLK)]], 26)
for j, (big, small, col) in enumerate([('3.2 mm', 'median hand error, from 90 mm', 'C25E00'), ('4.5°', 'whole-body rotation error', BLK), ('0.93', 'contact F1, learned support', BLK)]):
    text(s, 1392, 400 + j * 150, 400, 88, big, 80, col, HEAD, bold=True, spacing=1.0); text(s, 1392, 400 + j * 150 + 92, 400, 40, small, 26, '3D3B35')
footer(s, '600 held-out manipulation sequences · data 0.50M → 5.04M rows · model 3.7M → 57.7M parameters', '6D6A61')
notes(s, NOTES['loss'])

# tiangong
s = slide(BG); eyebrow(s, 'What this means · 4'); title(s, 'Every humanoid, its own whole-body controller', h=150)
video(s, f'{W}/tiangong-seven-action-reel', 128, 300, 1000, 563); video(s, f'{W}/tiangong-jump', 1152, 300, 640, 270); video(s, f'{W}/tiangong-flip', 1152, 593, 640, 270)
text(s, 128, 895, 1600, 50, 'Seven actions tracked in physics: turn and walk, dance, duck walk, squat, sit down, jump, flip.', 28, SOFT)
footer(s, 'Tiangong, zero-shot motion transfer in simulation · silver = robot, blue = reference')
notes(s, NOTES['refiner'])

# crawl
s = slide(BG); eyebrow(s, 'What this means · 4'); title(s, 'A Fourier GR3 learns to crawl from tokenizer data')
video(s, f'{W}/gmr-duel-crawl-01', 128, 300, 1664, 376)
for i, (b, sm, col) in enumerate([('4 / 4', 'crawl clips pass on tokenizer data (native retarget: 1 / 4)', ORANGE), ('≥ 96.9%', 'best success on 18 paired motions, one frozen controller', TXT), ('2.4×', 'smoother references (jerk)', TXT)]):
    x = 128 + i * 570; text(s, x, 730, 530, 88, b, 80, col, HEAD, bold=True, spacing=1.0); text(s, x, 824, 520, 80, sm, 26, MUTE, spacing=1.3)
footer(s, 'mjlab (MuJoCo-Warp) · BeyondMimic recipe · same controller, same budget'); notes(s, NOTES['physics'])

# open data
s = slide(ORANGE); eyebrow(s, 'Open data · IROS and CoRL', BLK)
text(s, 128, 200, 1100, 210, '500,000 h', 200, BLK, HEAD, bold=True, spacing=0.95)
text(s, 128, 420, 1040, 110, 'The largest cross-embodiment robot dataset to date. Free for academic and non-commercial use.', 40, BLK, spacing=1.25)
text(s, 128, 560, 1040, 90, 'Expanding to 20 embodiments: Unitree G1, H2 · AgiBot X2, A3 · Fourier GR3 · Booster T3.', 30, BLK, spacing=1.35)
text(s, 128, 680, 1040, 100, 'Hand us a new URDF, and all 500,000 hours map onto it.', 36, BLK, HEAD, bold=True, spacing=1.25)
text(s, 128, 810, 1040, 40, '128 GB200s running at full throttle.', 26, BLK)
image(s, f'{M}/qr_hf.png', 1312, 250, 420, 420)
text(s, 1282, 690, 480, 70, 'huggingface.co/datasets/sajio/alphamotion-soma-h2', 22, BLK, MONO, align=PP_ALIGN.CENTER)
footer(s, 'First release on Hugging Face today: 26,111 SOMA → Unitree H2 clips', '3D2A10'); notes(s, NOTES['open-data'])

# invite
s = slide(BG); eyebrow(s, 'An invitation'); title(s, 'Use the data. Use our compute. Both free.', size=72, h=90)
cw = (1664 - 48) / 2
for i, (a, b) in enumerate([('Data', '500,000 hours, free for academic and non-commercial use.'), ('Compute', 'Not enough compute to scale? Use our GB200s and H200s, free of charge.')]):
    x = 128 + i * (cw + 48); rect(s, x, 330, cw, 3, fill=ORANGE)
    text(s, x, 361, cw, 64, a, 48, TXT, HEAD, bold=True); text(s, x, 440, cw, 140, b, 32, SOFT, spacing=1.35)
text(s, 128, 700, 1500, 110, 'Next year, not only in Palo Alto: fully autonomous humanoids in every factory and on every street.', 34, MUTE, spacing=1.35)
footer(s, AUTH); notes(s, NOTES['invite'])

# thanks
s = slide(BG)
text(s, 128, 260, 1100, 120, 'The wheels to spin up the new era of embodied AI are in your hands.', 40, MUTE, spacing=1.35)
text(s, 128, 420, 1100, 220, 'Let’s turn the wheels together.', 96, TXT, HEAD, bold=True, spacing=1.05)
rect(s, 128, 680, 120, 3, fill=ORANGE)
text(s, 128, 710, 900, 50, 'Lloyd Lei · lloydl@andrew.cmu.edu', 32, '9C9A93')
image(s, f'{M}/qr_hf.png', 1392, 300, 400, 400)
text(s, 1392, 720, 400, 40, 'Open dataset', 24, '9C9A93', align=PP_ALIGN.CENTER)
footer(s, AUTH); notes(s, NOTES['thanks'])

from lxml import etree
NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
def autoplay(sl):
    ids = [sh.shape_id for sh in sl.shapes if sh.shape_type is not None and sh.shape_type == 16]
    if not ids: return 0
    old = sl.element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}timing')
    if old is not None: sl.element.remove(old)
    n = [5]
    def nid():
        n[0] += 1; return n[0]
    calls = ''.join(f'<p:par><p:cTn id="{nid()}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="withEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{i}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>' for i in ids)
    media = ''.join(f'<p:video><p:cMediaNode vol="0" mute="1"><p:cTn id="{nid()}" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{i}"/></p:tgtEl></p:cMediaNode></p:video>' for i in ids)
    xml = (f'<p:timing {NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
           f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{calls}</p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'{media}</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    el = etree.fromstring(xml)
    ext = sl.element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}extLst')
    if ext is not None: ext.addprevious(el)
    else: sl.element.append(el)
    return len(ids)
print('autoplay videos per slide:', [autoplay(sl) for sl in prs.slides])
prs.save(OUT); print('saved', OUT, os.path.getsize(OUT) // 1024, 'KB', len(prs.slides), 'slides')
